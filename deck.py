"""Taqdimot dvigateli (Gamma/Canva uslubida).

Taqdimot FAYL emas, balki TUZILMA (JSON) sifatida saqlanadi: slaydlar ro'yxati, har birida layout va maydonlar.
Shu tufayli «3-slaydni qisqartir» deyilganda faqat shu slayd o'zgaradi, qolganlari buzilmaydi.
build_pptx() tuzilmadan chiroyli .pptx yasaydi (PowerPoint, Google Slides, Canva ochadi)."""
import copy
import io
import json
import math
import re

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


class DeckError(Exception):
    """Foydalanuvchiga ko'rsatish mumkin bo'lgan xato."""

    def __init__(self, message: str, *, kind: str = "invalid_deck", detail: str = ""):
        super().__init__(message)
        self.kind = kind
        self.detail = detail


W, H = 13.333, 7.5
MX = 0.8
CW = W - 2 * MX
FONT = "Calibri"
MAX_SLIDES = 20
MIN_SLIDES = 5  # YANGI: minimal slayd soni

LAYOUTS = ("title", "section", "bullets", "image_text", "two_column", "stats",
           "steps", "quote", "table", "chart", "closing")

# bg/bg2 — fon; card — kartochka; line — chiziq; hero1/hero2 — sarlavha slaydlari gradienti
THEMES = {
    "midnight": dict(dark=True, bg="0F172A", bg2="1E293B", card="1E293B", line="334155", text="F8FAFC",
                     muted="94A3B8", accent="38BDF8", accent2="818CF8", on_accent="0F172A",
                     hero1="0B1220", hero2="1E3A5F", hero_text="FFFFFF", hero_sub="BFD3EA"),
    "ocean": dict(dark=False, bg="F0F9FF", bg2="E0F2FE", card="FFFFFF", line="BAE6FD", text="0C4A6E",
                  muted="475569", accent="0284C7", accent2="06B6D4", on_accent="FFFFFF",
                  hero1="075985", hero2="06B6D4", hero_text="FFFFFF", hero_sub="E0F7FF"),
    "sunset": dict(dark=True, bg="2A1B3D", bg2="44174E", card="3B2150", line="5B3A73", text="FFF7ED",
                   muted="D8B4FE", accent="FB7185", accent2="FBBF24", on_accent="2A1B3D",
                   hero1="3B0F4F", hero2="C2410C", hero_text="FFFFFF", hero_sub="FFE4D0"),
    "forest": dict(dark=False, bg="F7FBF4", bg2="E3F1DC", card="FFFFFF", line="CDE5C2", text="14532D",
                   muted="4B5563", accent="16A34A", accent2="CA8A04", on_accent="FFFFFF",
                   hero1="14532D", hero2="4D9B5F", hero_text="FFFFFF", hero_sub="E3F5DD"),
    "royal": dict(dark=True, bg="1E1B4B", bg2="312E81", card="2B2766", line="4C47A3", text="EEF2FF",
                  muted="C4B5FD", accent="A78BFA", accent2="F472B6", on_accent="1E1B4B",
                  hero1="1E1B4B", hero2="6D28D9", hero_text="FFFFFF", hero_sub="DDD6FE"),
    "minimal": dict(dark=False, bg="FFFFFF", bg2="F4F4F5", card="F4F4F5", line="E4E4E7", text="18181B",
                    muted="52525B", accent="E11D48", accent2="18181B", on_accent="FFFFFF",
                    hero1="18181B", hero2="3F3F46", hero_text="FFFFFF", hero_sub="D4D4D8"),
}
DEFAULT_THEME = "midnight"
CHART_EXTRA = ["F59E0B", "10B981", "EF4444", "8B5CF6", "14B8A6", "F97316"]


# ================================================================ JSON → tuzilma
def _s(v, n: int) -> str:
    s = re.sub(r"\s+", " ", str(v if v is not None else "")).strip()
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


def _list(v, n_items: int, n_chars: int) -> list[str]:
    if isinstance(v, str):
        v = [v]
    if not isinstance(v, list):
        return []
    return [_s(x, n_chars) for x in v if _s(x, n_chars)][:n_items]


def _num(v):
    try:
        if isinstance(v, str):
            v = v.replace(" ", "").replace(",", ".").rstrip("%")
        return float(v)
    except (TypeError, ValueError):
        return None


def _norm_slide(s) -> dict | None:
    if not isinstance(s, dict):
        return None
    lay = str(s.get("layout") or "bullets").strip().lower()
    if lay not in LAYOUTS:
        lay = "bullets"
    out: dict = {"layout": lay}
    title = _s(s.get("title"), 90)
    sub = _s(s.get("subtitle"), 160)
    bullets = _list(s.get("bullets"), 7, 150)
    img = _s(s.get("image_prompt"), 240)
    if title:
        out["title"] = title
    if sub:
        out["subtitle"] = sub
    if s.get("notes"):
        out["notes"] = _s(s.get("notes"), 600)

    if lay == "two_column":
        left, right = _list(s.get("left"), 5, 120), _list(s.get("right"), 5, 120)
        if left and right:
            out.update(left=left, right=right,
                       left_title=_s(s.get("left_title"), 40), right_title=_s(s.get("right_title"), 40))
        else:
            bullets, lay = bullets or (left + right), "bullets"
    elif lay == "stats":
        stats = []
        for it in (s.get("stats") if isinstance(s.get("stats"), list) else []):
            if isinstance(it, dict) and _s(it.get("value"), 14) and _s(it.get("label"), 70):
                stats.append({"value": _s(it["value"], 14), "label": _s(it["label"], 70)})
        if len(stats) >= 2:
            out["stats"] = stats[:4]
        else:
            lay = "bullets"
    elif lay == "steps":
        steps = []
        for it in (s.get("steps") if isinstance(s.get("steps"), list) else []):
            if isinstance(it, dict) and _s(it.get("title"), 40):
                steps.append({"title": _s(it["title"], 40), "text": _s(it.get("text"), 110)})
            elif isinstance(it, str) and _s(it, 40):
                steps.append({"title": _s(it, 40), "text": ""})
        if len(steps) >= 2:
            out["steps"] = steps[:5]
        else:
            lay = "bullets"
    elif lay == "quote":
        if _s(s.get("quote"), 300):
            out["quote"] = _s(s.get("quote"), 300)
            out["author"] = _s(s.get("author"), 80)
        else:
            lay = "bullets"
    elif lay == "table":
        headers = _list(s.get("headers"), 5, 40)
        rows = []
        for r in (s.get("rows") if isinstance(s.get("rows"), list) else []):
            if isinstance(r, list):
                rows.append([_s(c, 60) for c in r][: len(headers)])
        rows = [r + [""] * (len(headers) - len(r)) for r in rows][:7]
        if len(headers) >= 2 and rows:
            out.update(headers=headers, rows=rows)
        else:
            lay = "bullets"
    elif lay == "chart":
        ch = s.get("chart") if isinstance(s.get("chart"), dict) else {}
        cats = _list(ch.get("categories"), 8, 24)
        series = []
        for se in (ch.get("series") if isinstance(ch.get("series"), list) else []):
            if not isinstance(se, dict) or not isinstance(se.get("values"), list):
                continue
            vals = [_num(v) for v in se["values"]][: len(cats)]
            if len(vals) == len(cats) and all(v is not None for v in vals):
                series.append({"name": _s(se.get("name"), 30) or "Qiymat", "values": vals})
        ctype = str(ch.get("type") or "bar").lower()
        if ctype not in ("bar", "line", "pie", "barh"):
            ctype = "bar"
        if len(cats) >= 2 and series:
            out["chart"] = {"type": ctype, "categories": cats, "series": series[:3]}
        else:
            lay = "bullets"
    elif lay == "image_text":
        if not (bullets and img):
            lay = "bullets"

    out["layout"] = lay
    if lay in ("bullets", "image_text"):
        if not bullets:
            out["layout"] = "section"
        else:
            out["bullets"] = bullets[:6]
    elif lay == "chart":
        if bullets:
            out["bullets"] = bullets[:3]
    if img and lay in ("title", "image_text", "section", "closing"):
        out["image_prompt"] = img
    if out["layout"] in ("bullets", "image_text", "two_column", "stats", "steps", "table", "chart") \
            and "title" not in out:
        out["title"] = "Slayd"
    if out["layout"] in ("title", "section", "closing") and "title" not in out:
        out["title"] = "Taqdimot"
    return out


def normalize(d) -> dict:
    _validate_raw_deck(d)
    slides = [x for x in (_norm_slide(s) for s in d["slides"]) if x][:MAX_SLIDES]
    if not slides:
        raise DeckError("Taqdimotda slaydlar topilmadi.")
    if len(slides) < MIN_SLIDES:
        raise DeckError(
            f"Taqdimot juda qisqa ({len(slides)} slayd). "
            f"Kamida {MIN_SLIDES} ta kerak. Qayta urinib ko'ring."
        )
    # Har bir mazmunli matnli slaydni rasmli layoutga o'tkazamiz. Rasm API
    # ishlamasa ham build_pptx shu layoutni xilma-xil matn ko'rinishida chizadi.
    for index, s in enumerate(slides):
        if s.get("layout") == "title" and not s.get("subtitle"):
            s["subtitle"] = "Mavzuga kirish va taqdimotning asosiy maqsadi"
        elif s.get("layout") == "section" and not s.get("subtitle"):
            s["subtitle"] = "Bo'limning asosiy tushunchalari, izohlari va amaliy misollari"
        elif s.get("layout") == "closing" and not s.get("subtitle"):
            s["subtitle"] = "Asosiy fikrlar, amaliy xulosa va eslab qolish kerak bo'lgan jihatlar"
        if not s.get("image_prompt"):
            title = s.get("title") or s.get("quote", "")[:40] or "presentation"
            s["image_prompt"] = (
                f"Create a clear, high-quality editorial visual for slide {index + 1} "
                f"about {title}. Make this scene and camera composition distinct "
                "from other slides in the same deck; modern, detailed, 16:9."
            )
        if s.get("layout") == "bullets":
            s["layout"] = "image_text"
    theme = str(d.get("theme") or DEFAULT_THEME).strip().lower()
    return {
        "title": _s(d.get("title") or slides[0].get("title"), 90) or "Taqdimot",
        "theme": theme if theme in THEMES else DEFAULT_THEME,
        "slides": slides,
    }


def _validate_raw_deck(deck) -> None:
    """Reject malformed provider output before normalization can hide its shape."""
    if not isinstance(deck, dict):
        raise DeckError("DeckError: javobning yuqori darajasi JSON obyekt bo'lishi kerak.")
    for field in ("title", "theme"):
        if field in deck and deck[field] is not None and not isinstance(deck[field], str):
            raise DeckError(f"DeckError: '{field}' maydoni matn bo'lishi kerak.")
    slides = deck.get("slides")
    if not isinstance(slides, list):
        raise DeckError("DeckError: 'slides' maydoni JSON ro'yxati bo'lishi kerak.")
    if not slides:
        raise DeckError("DeckError: 'slides' ro'yxati bo'sh.")
    if len(slides) > MAX_SLIDES:
        raise DeckError(f"DeckError: {len(slides)} slayd qaytdi, maksimum {MAX_SLIDES}.")

    string_fields = ("title", "subtitle", "image_prompt", "notes", "quote", "author",
                     "left_title", "right_title")
    list_fields = ("bullets", "left", "right", "headers", "rows", "steps", "stats")
    for index, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            raise DeckError(f"DeckError: {index}-slayd JSON obyekt emas ({type(slide).__name__}).")
        for field in string_fields:
            if field in slide and slide[field] is not None and not isinstance(slide[field], str):
                raise DeckError(f"DeckError: {index}-slayddagi '{field}' matn bo'lishi kerak.")
        for field in list_fields:
            if field in slide and slide[field] is not None and not isinstance(slide[field], (list, str)):
                raise DeckError(f"DeckError: {index}-slayddagi '{field}' ro'yxat yoki matn bo'lishi kerak.")
        for field in ("bullets", "left", "right", "headers"):
            value = slide.get(field)
            if isinstance(value, list) and any(not isinstance(item, str) for item in value):
                raise DeckError(f"DeckError: {index}-slayddagi '{field}' elementlari matn bo'lishi kerak.")
        if isinstance(slide.get("rows"), list) and any(
            not isinstance(row, list) or any(not isinstance(cell, (str, int, float)) for cell in row)
            for row in slide["rows"]
        ):
            raise DeckError(f"DeckError: {index}-slayddagi 'rows' qatorlar ro'yxati bo'lishi kerak.")
        if isinstance(slide.get("stats"), list) and any(not isinstance(item, dict) for item in slide["stats"]):
            raise DeckError(f"DeckError: {index}-slayddagi 'stats' obyektlar ro'yxati bo'lishi kerak.")
        if isinstance(slide.get("steps"), list) and any(not isinstance(item, (dict, str)) for item in slide["steps"]):
            raise DeckError(f"DeckError: {index}-slayddagi 'steps' obyekt yoki matn ro'yxati bo'lishi kerak.")
        layout = slide.get("layout")
        if layout is not None and (not isinstance(layout, str) or layout.strip().lower() not in LAYOUTS):
            raise DeckError(f"DeckError: {index}-slayd layout qiymati tanilmagan.")
        if "chart" in slide and slide["chart"] is not None and not isinstance(slide["chart"], dict):
            raise DeckError(f"DeckError: {index}-slayddagi 'chart' obyekt bo'lishi kerak.")
        if layout == "two_column" and not all(slide.get(key) for key in ("left", "right")):
            raise DeckError(f"DeckError: {index}-slayd two_column uchun left va right talab qilinadi.")
        if layout in ("bullets", "image_text") and not slide.get("bullets"):
            raise DeckError(f"DeckError: {index}-slaydda bullets matni bo'sh.")
        if layout == "stats" and not isinstance(slide.get("stats"), list):
            raise DeckError(f"DeckError: {index}-slaydda stats ro'yxati talab qilinadi.")
        if layout == "stats" and len(slide["stats"]) < 2:
            raise DeckError(f"DeckError: {index}-slaydda kamida ikki stats elementi kerak.")
        if layout == "steps" and not isinstance(slide.get("steps"), list):
            raise DeckError(f"DeckError: {index}-slaydda steps ro'yxati talab qilinadi.")
        if layout == "steps" and len(slide["steps"]) < 2:
            raise DeckError(f"DeckError: {index}-slaydda kamida ikki steps elementi kerak.")
        if layout == "quote" and not slide.get("quote"):
            raise DeckError(f"DeckError: {index}-slaydda quote matni talab qilinadi.")
        if layout == "table" and not all(slide.get(key) for key in ("headers", "rows")):
            raise DeckError(f"DeckError: {index}-slayd table uchun headers va rows talab qilinadi.")
        if layout == "chart" and not isinstance(slide.get("chart"), dict):
            raise DeckError(f"DeckError: {index}-slaydda chart obyekt talab qilinadi.")
        if layout in ("title", "section", "closing") and not (slide.get("title") or slide.get("subtitle")):
            raise DeckError(f"DeckError: {index}-slayd uchun title yoki subtitle talab qilinadi.")


def parse_json(raw: str) -> dict:
    """Model javobidan JSON ajratadi (``` va ortiqcha matnni olib tashlaydi)."""
    text = re.sub(r"^```[a-zA-Z]*\s*|\s*```$", "", (raw or "").strip())
    a, b = text.find("{"), text.rfind("}")
    if a < 0 or b < a:
        raise DeckError(
            "Model taqdimotni JSON ko'rinishida qaytarmadi, qayta urinib ko'ring.",
            kind="missing_json_object", detail="No JSON object delimiters found",
        )
    try:
        return json.loads(text[a:b + 1])
    except json.JSONDecodeError as exc:
        detail = f"line={exc.lineno}; column={exc.colno}; position={exc.pos}; reason={exc.msg[:120]}"
        raise DeckError("Model javobi buzilgan (JSON), qayta urinib ko'ring.",
                        kind="invalid_json", detail=detail) from exc


def compact(deck: dict) -> str:
    """Modelga berish uchun: slaydlar 1 dan raqamlangan JSON."""
    view = {"title": deck["title"], "theme": deck["theme"],
            "slides": [{"n": i + 1, **s} for i, s in enumerate(deck["slides"])]}
    return json.dumps(view, ensure_ascii=False)


def outline(deck: dict) -> str:
    names = {"title": "sarlavha", "section": "bo'lim", "bullets": "matn", "image_text": "rasmli",
             "two_column": "2 ustun", "stats": "raqamlar", "steps": "bosqichlar", "quote": "iqtibos",
             "table": "jadval", "chart": "diagramma", "closing": "yakun"}
    rows = []
    for i, s in enumerate(deck["slides"], 1):
        t = s.get("title") or s.get("quote", "")[:40]
        rows.append(f"{i}. {t} ({names.get(s['layout'], s['layout'])})")
    return "\n".join(rows)


def image_prompts(deck: dict) -> list[str]:
    seen, out = set(), []
    for s in deck["slides"]:
        if s.get("layout") not in ("title", "image_text", "closing"):
            continue
        p = s.get("image_prompt")
        if p and p not in seen:
            seen.add(p)
            out.append(p)
    return out


# ================================================================ tahrir amallari
def apply_ops(deck: dict, ops) -> tuple[dict, list[str]]:
    """Model qaytargan amallarni qo'llaydi. (yangi taqdimot, o'zgarishlar ro'yxati)."""
    d = copy.deepcopy(deck)
    slides = d["slides"]
    done: list[str] = []
    if not isinstance(ops, list):
        ops = []
    for op in ops:
        if not isinstance(op, dict):
            continue
        kind = str(op.get("op") or "").lower()
        try:
            n = int(op.get("slide", 0))
        except (TypeError, ValueError):
            n = 0
        if kind == "replace" and 1 <= n <= len(slides) and isinstance(op.get("content"), dict):
            old, new = slides[n - 1], op["content"]
            fixed = _norm_slide({**old, **new})
            if fixed:
                slides[n - 1] = fixed
                done.append(f"{n}-slayd o'zgartirildi")
        elif kind == "insert" and isinstance(op.get("content"), dict) and len(slides) < MAX_SLIDES:
            try:
                after = int(op.get("after", len(slides)))
            except (TypeError, ValueError):
                after = len(slides)
            after = min(max(after, 0), len(slides))
            fixed = _norm_slide(op["content"])
            if fixed:
                slides.insert(after, fixed)
                done.append(f"{after + 1}-o'ringa yangi slayd qo'shildi")
        elif kind == "delete" and 1 <= n <= len(slides) and len(slides) > 1:
            slides.pop(n - 1)
            done.append(f"{n}-slayd o'chirildi")
        elif kind == "move" and 1 <= n <= len(slides):
            try:
                to = min(max(int(op.get("to", n)), 1), len(slides))
            except (TypeError, ValueError):
                to = n
            if to != n:
                slides.insert(to - 1, slides.pop(n - 1))
                done.append(f"{n}-slayd {to}-o'ringa ko'chirildi")
        elif kind == "theme":
            name = str(op.get("name") or "").strip().lower()
            if name in THEMES and name != d["theme"]:
                d["theme"] = name
                done.append(f"mavzu: {name}")
        elif kind == "title" and _s(op.get("text"), 90):
            d["title"] = _s(op.get("text"), 90)
            done.append("taqdimot nomi o'zgartirildi")
    if not done:
        raise DeckError("Qaysi slaydni qanday o'zgartirish kerakligini tushunmadim. "
                        "Masalan: «3-slaydni qisqartir» yoki «mavzuni ocean qil».")
    return d, done


# ================================================================ chizish yordamchilari
def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _alpha(shape, opacity_pct: float) -> None:
    sf = shape._element.spPr.find(qn("a:solidFill"))
    if sf is None or not len(sf):
        return
    clr = sf[0]
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(opacity_pct * 1000)))


def _shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None, opacity=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        if opacity is not None:
            _alpha(shp, opacity)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def rect(slide, x, y, w, h, fill, **kw):
    return _shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, **kw)


def rrect(slide, x, y, w, h, fill, radius=0.08, **kw):
    return _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=radius, **kw)


def oval(slide, x, y, w, h, fill, **kw):
    return _shape(slide, MSO_SHAPE.OVAL, x, y, w, h, fill, **kw)


def _lines(text: str, size: float, w: float, bold=False) -> int:
    cpl = max(1, int(w * 72 / (size * (0.55 if bold else 0.5))))
    return sum(max(1, math.ceil(len(p) / cpl)) for p in str(text).split("\n"))


def fit(texts, w: float, h: float, max_s: int, min_s: int, bold=False, gap=0.08, lh=1.22) -> int:
    texts = [t for t in texts if t]
    if not texts:
        return min_s
    s = max_s
    while s > min_s:
        total = sum(_lines(t, s, w, bold) * s * lh / 72 for t in texts) + gap * max(len(texts) - 1, 0)
        if total <= h:
            break
        s -= 1
    return s


def text(slide, x, y, w, h, paras, size, color, bold=False, italic=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=0):
    """paras: str | list[str | list[(text, bold, size|None)]]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = [(para, bold, None)] if isinstance(para, str) else para
        for t, b, sz in runs:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(sz or size)
            r.font.bold = b
            r.font.italic = italic
            r.font.color.rgb = _rgb(color)
    return tb


def _split_head(b: str) -> tuple[str, str]:
    for sep in (": ", " — ", " – "):
        if sep in b:
            head, body = b.split(sep, 1)
            if 2 <= len(head) <= 42 and body.strip():
                return head.strip(), body.strip()
    return "", b


def _crop(img: bytes, w_in: float, h_in: float) -> bytes:
    from PIL import Image

    if w_in <= 0 or h_in <= 0:
        raise ValueError("Rasm o'lchamlari noto'g'ri")
    im = Image.open(io.BytesIO(img)).convert("RGB")
    tw, th = im.size
    if tw == 0 or th == 0:
        raise ValueError("Rasm bo'sh")
    target = w_in / h_in
    if tw / th > target:
        nw = int(th * target)
        im = im.crop(((tw - nw) // 2, 0, (tw - nw) // 2 + nw, th))
    else:
        nh = int(tw / target)
        im = im.crop((0, (th - nh) // 2, tw, (th - nh) // 2 + nh))
    px = min(int(w_in * 140), 1800)
    im = im.resize((px, max(1, int(px / target))))
    out = io.BytesIO()
    im.save(out, "JPEG", quality=86)
    return out.getvalue()


def _picture(slide, img: bytes, x, y, w, h) -> bool:
    try:
        slide.shapes.add_picture(io.BytesIO(_crop(img, w, h)), Inches(x), Inches(y), Inches(w), Inches(h))
        return True
    except Exception:  # noqa: BLE001 — buzilgan rasm butun taqdimotni to'xtatmasin
        return False


def _hero_bg(slide, th) -> None:
    f = slide.background.fill
    f.gradient()
    f.gradient_angle = 35
    f.gradient_stops[0].color.rgb = _rgb(th["hero1"])
    f.gradient_stops[1].color.rgb = _rgb(th["hero2"])


def _solid_bg(slide, color: str) -> None:
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = _rgb(color)


def _header(slide, th, title: str) -> None:
    rect(slide, MX, 0.62, 0.9, 0.07, th["accent"])
    size = fit([title], CW, 1.0, 36, 22, bold=True, lh=1.12)
    text(slide, MX, 0.82, CW, 1.05, title, size, th["text"], bold=True)


def _footer(slide, th, deck_title: str, num: int) -> None:
    text(slide, MX, 7.02, 8, 0.25, deck_title, 10, th["muted"])
    text(slide, W - MX - 1.2, 7.02, 1.2, 0.25, str(num), 10, th["muted"], align=PP_ALIGN.RIGHT)


def _dots(slide, th, items, x, y, w, h, max_s=22) -> None:
    if not items:
        return
    n = len(items)
    row = h / n
    size = min(fit([b], w - 0.5, row - 0.1, max_s, 13) for b in items)
    for i, b in enumerate(items):
        top = y + i * row
        head, body = _split_head(b)
        oval(slide, x, top + size / 72 * 0.45, 0.17, 0.17, th["accent"])
        paras = [[(head + ": ", True, None), (body, False, None)]] if head else [b]
        text(slide, x + 0.4, top, w - 0.4, row - 0.08, paras, size, th["text"])


# ================================================================ slayd turlari
def _slide_title(slide, th, s, img, deck_title, num, closing=False):
    _hero_bg(slide, th)
    big = oval(slide, W - 5.2, -2.3, 7.6, 7.6, th["accent"], opacity=16)
    oval(slide, -2.2, H - 3.2, 5.4, 5.4, th["accent2"], opacity=14)
    del big
    has_img = bool(img) and not closing
    if has_img:
        shown = _picture(slide, img, 8.25, 0, W - 8.25, H)
        has_img = shown
    tw = 6.9 if has_img else (9.6 if not closing else 10.6)
    x = MX + 0.1 if not closing else (W - tw) / 2
    al = PP_ALIGN.LEFT if not closing else PP_ALIGN.CENTER
    title, sub = s.get("title", ""), s.get("subtitle", "")
    size = fit([title], tw, 2.6, 56, 30, bold=True, lh=1.1)
    lines = _lines(title, size, tw, True)
    th_h = lines * size * 1.14 / 72
    top = max(1.4, (H - th_h - (1.3 if sub else 0)) / 2 - 0.2)
    bar_x = x if not closing else (W - 1.1) / 2
    rect(slide, bar_x, top - 0.35, 1.1, 0.09, th["accent2"] if th["dark"] else "FFFFFF")
    text(slide, x, top, tw, th_h + 0.1, title, size, th["hero_text"], bold=True, align=al)
    if sub:
        ssz = fit([sub], tw, 1.3, 24, 14)
        text(slide, x, top + th_h + 0.3, tw, 1.3, sub, ssz, th["hero_sub"], align=al)


def _slide_section(slide, th, s, num, idx):
    _hero_bg(slide, th)
    oval(slide, W - 4.4, H - 4.6, 7, 7, th["accent2"], opacity=15)
    text(slide, MX + 0.1, 1.5, 6, 2.2, f"{idx:02d}", 120, th["hero_sub"], bold=True)
    title, sub = s.get("title", ""), s.get("subtitle", "")
    size = fit([title], 10.5, 1.8, 48, 28, bold=True, lh=1.1)
    rect(slide, MX + 0.1, 4.0, 1.1, 0.09, th["accent2"] if th["dark"] else "FFFFFF")
    text(slide, MX + 0.1, 4.3, 10.5, 1.8, title, size, th["hero_text"], bold=True)
    if sub:
        text(slide, MX + 0.1, 6.1, 10.5, 0.8, sub, 20, th["hero_sub"])


def _slide_bullets(slide, th, s, variant=0):
    items = s.get("bullets", [])
    if not items:
        return
    if variant == 1:
        row_h = 4.7 / len(items)
        for index, item in enumerate(items):
            y = 2.05 + index * row_h
            if index % 2 == 0:
                rrect(slide, MX, y, CW, row_h - 0.1, th["card"], line=th["line"])
            oval(slide, MX + 0.2, y + (row_h - 0.46) / 2, 0.46, 0.46, th["accent"])
            text(slide, MX + 0.2, y + (row_h - 0.46) / 2, 0.46, 0.46,
                 str(index + 1), 16, th["on_accent"], bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            head, body = _split_head(item)
            paras = [[(head + ": ", True, None), (body, False, None)]] if head else [item]
            size = fit([item], CW - 1.1, row_h - 0.16, 21, 12)
            text(slide, MX + 0.9, y + 0.04, CW - 1.1, row_h - 0.16,
                 paras, size, th["text"], anchor=MSO_ANCHOR.MIDDLE)
        return
    if variant == 2 and len(items) > 1:
        hero_w = min(4.2, CW * 0.4)
        rrect(slide, MX, 2.05, hero_w, 4.7, th["accent"], radius=0.06)
        head, body = _split_head(items[0])
        lead = f"{head}: {body}" if head else items[0]
        lead_size = fit([lead], hero_w - 0.55, 4.0, 30, 16, bold=True)
        text(slide, MX + 0.28, 2.35, hero_w - 0.55, 4.0, lead,
             lead_size, th["on_accent"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _dots(slide, th, items[1:], MX + hero_w + 0.35, 2.15,
              CW - hero_w - 0.45, 4.5, 21)
        return
    n = len(items)
    cols = 1 if n <= 3 else 2
    rows = math.ceil(n / cols)
    gap = 0.28
    area_y, area_h = 2.05, 4.75
    cw = (CW - (cols - 1) * gap) / cols
    ch = min((area_h - (rows - 1) * gap) / rows, 2.1)
    top = area_y + (area_h - (rows * ch + (rows - 1) * gap)) / 2
    for i, b in enumerate(items):
        r, c = divmod(i, cols)
        x, y = MX + c * (cw + gap), top + r * (ch + gap)
        rrect(slide, x, y, cw, ch, th["card"], line=th["line"])
        d = 0.58
        oval(slide, x + 0.28, y + (ch - d) / 2, d, d, th["accent"])
        text(slide, x + 0.28, y + (ch - d) / 2, d, d, str(i + 1), 20, th["on_accent"], bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tx, tw = x + 1.15, cw - 1.45
        head, body = _split_head(b)
        if head:
            size = fit([head, body], tw, ch - 0.3, 22, 12)
            paras = [[(head, True, size + 1)], [(body, False, size - 1)]]
        else:
            size = fit([b], tw, ch - 0.3, 24, 13)
            paras = [b]
        text(slide, tx, y + 0.15, tw, ch - 0.3, paras, size, th["text"], anchor=MSO_ANCHOR.MIDDLE,
             space_after=2)


def _slide_image_text(slide, th, s, img, variant=0):
    px = 7.95
    bullets = s.get("bullets", [])
    if img and bullets and _picture(slide, img, px, 0, W - px, H):
        rect(slide, px - 0.07, 0, 0.07, H, th["accent"])
        rect(slide, MX, 0.62, 0.9, 0.07, th["accent"])
        title = s.get("title", "Slayd")
        size = fit([title], px - MX - 0.6, 1.5, 34, 22, bold=True, lh=1.12)
        text(slide, MX, 0.82, px - MX - 0.6, 1.5, title, size, th["text"], bold=True)
        content_w = px - MX - 0.7
        if variant == 1:
            gap = 0.12
            row_h = 4.1 / max(1, len(bullets))
            for index, item in enumerate(bullets):
                y = 2.5 + index * row_h
                if index % 2 == 0:
                    rrect(slide, MX, y, content_w, row_h - gap, th["card"], line=th["line"])
                text(slide, MX + 0.12, y + 0.04, 0.42, row_h - gap - 0.08,
                     f"{index + 1:02d}", 14, th["accent"], bold=True,
                     anchor=MSO_ANCHOR.MIDDLE)
                head, body = _split_head(item)
                paras = [[(head + ": ", True, None), (body, False, None)]] if head else [item]
                size = fit([item], content_w - 0.8, row_h - gap - 0.08, 19, 11)
                text(slide, MX + 0.62, y + 0.04, content_w - 0.8,
                     row_h - gap - 0.08, paras, size, th["text"],
                     anchor=MSO_ANCHOR.MIDDLE)
        elif variant == 2 and len(bullets) > 1:
            hero_w = content_w * 0.43
            rrect(slide, MX, 2.5, hero_w, 4.1, th["accent"], radius=0.06)
            head, body = _split_head(bullets[0])
            lead = f"{head}: {body}" if head else bullets[0]
            lead_size = fit([lead], hero_w - 0.55, 3.5, 28, 15, bold=True)
            text(slide, MX + 0.28, 2.8, hero_w - 0.55, 3.5, lead,
                 lead_size, th["on_accent"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
            _dots(slide, th, bullets[1:], MX + hero_w + 0.2, 2.55,
                  content_w - hero_w - 0.2, 4.0, 19)
        else:
            _dots(slide, th, bullets, MX, 2.5, content_w, 4.2, 22)
    else:
        _header(slide, th, s.get("title", "Slayd"))
        _slide_bullets(slide, th, s, variant)


def _slide_two_column(slide, th, s):
    left = s.get("left", [])
    right = s.get("right", [])
    if not left and not right:
        return
    gap = 0.4
    cw = (CW - gap) / 2
    for i, (ttl, items) in enumerate(((s.get("left_title", ""), left), (s.get("right_title", ""), right))):
        x = MX + i * (cw + gap)
        rrect(slide, x, 2.0, cw, 4.75, th["card"], line=th["line"], radius=0.05)
        y = 2.3
        if ttl:
            text(slide, x + 0.4, y, cw - 0.8, 0.6, ttl, 24, th["accent"], bold=True)
            rect(slide, x + 0.4, y + 0.62, 0.7, 0.06, th["accent"])
            y += 0.95
        if items:
            _dots(slide, th, items, x + 0.4, y, cw - 0.8, max(0.5, 6.55 - y), 20)


def _slide_stats(slide, th, s):
    stats = s.get("stats", [])
    if not stats:
        return
    n = len(stats)
    gap = 0.35
    cw = (CW - (n - 1) * gap) / n
    for i, st in enumerate(stats):
        x = MX + i * (cw + gap)
        rrect(slide, x, 2.25, cw, 3.55, th["card"], line=th["line"], radius=0.06)
        rect(slide, x + (cw - 0.8) / 2, 2.25, 0.8, 0.09, th["accent"])
        vs = fit([st["value"]], cw - 0.5, 1.4, 66, 28, bold=True, lh=1.0)
        text(slide, x + 0.25, 2.6, cw - 0.5, 1.45, st["value"], vs, th["accent"], bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        ls = fit([st["label"]], cw - 0.6, 1.5, 24, 14)
        text(slide, x + 0.3, 4.15, cw - 0.6, 1.5, st["label"], ls, th["text"], align=PP_ALIGN.CENTER)
    if s.get("subtitle"):
        text(slide, MX, 6.2, CW, 0.5, s["subtitle"], 16, th["muted"], align=PP_ALIGN.CENTER)


def _slide_steps(slide, th, s):
    steps = s.get("steps", [])
    if not steps:
        return
    n = len(steps)
    cw = CW / n
    rect(slide, MX + cw / 2, 2.87, CW - cw, 0.05, th["line"])
    for i, st in enumerate(steps):
        cx = MX + (i + 0.5) * cw
        oval(slide, cx - 0.4, 2.5, 0.8, 0.8, th["accent"])
        text(slide, cx - 0.4, 2.5, 0.8, 0.8, str(i + 1), 26, th["on_accent"], bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        ts = fit([st["title"]], cw - 0.3, 1.0, 22, 14, bold=True)
        text(slide, cx - cw / 2 + 0.15, 3.6, cw - 0.3, 1.0, st["title"], ts, th["text"], bold=True,
             align=PP_ALIGN.CENTER)
        if st.get("text"):
            bs = fit([st["text"]], cw - 0.4, 2.2, 18, 12)
            text(slide, cx - cw / 2 + 0.2, 4.55, cw - 0.4, 2.2, st["text"], bs, th["muted"],
                 align=PP_ALIGN.CENTER)


def _slide_quote(slide, th, s):
    if not s.get("quote"):
        return
    text(slide, MX, 0.6, 3, 2.4, "“", 160, th["accent"], bold=True)
    q = s["quote"]
    size = fit([q], 10.4, 3.2, 40, 20, lh=1.25)
    text(slide, 1.7, 2.1, 10.4, 3.4, q, size, th["text"], italic=True, anchor=MSO_ANCHOR.MIDDLE)
    if s.get("author"):
        rect(slide, 1.7, 5.85, 0.7, 0.06, th["accent"])
        text(slide, 1.7, 6.0, 10, 0.5, s["author"], 20, th["muted"], bold=True)


def _slide_table(slide, th, s):
    headers = s.get("headers", [])
    rows = s.get("rows", [])
    if not headers or not rows:
        return
    nr, nc = len(rows) + 1, len(headers)
    x, y, w = MX, 2.1, CW
    h = min(0.92 * nr, 4.7)
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gf.table
    weights = [max(6, max(len(str(r[c])) for r in [headers] + rows)) for c in range(nc)]
    tot = sum(weights) or 1
    size = 22 if nr <= 5 and nc <= 3 else (18 if nr <= 6 else 15)
    for c in range(nc):
        tbl.columns[c].width = Inches(w * weights[c] / tot)
    for r in range(nr):
        tbl.rows[r].height = Inches(h / nr)
        for c in range(nc):
            cell = tbl.cell(r, c)
            val = headers[c] if r == 0 else rows[r - 1][c]
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(th["accent"] if r == 0 else (th["card"] if r % 2 else th["bg2"]))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.15)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            run = p.add_run()
            run.text = val
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = r == 0
            run.font.color.rgb = _rgb(th["on_accent"] if r == 0 else th["text"])


def _slide_chart(slide, th, s):
    ch = s.get("chart")
    if not ch or not ch.get("categories") or not ch.get("series"):
        return
    kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "barh": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[ch["type"]]
    data = CategoryChartData()
    data.categories = ch["categories"]
    series = ch["series"][:1] if ch["type"] == "pie" else ch["series"]
    for se in series:
        data.add_series(se["name"], se["values"])
    side = bool(s.get("bullets"))
    cw = 7.7 if side else CW
    gf = slide.shapes.add_chart(kind, Inches(MX), Inches(2.0), Inches(cw), Inches(4.85), data)
    c = gf.chart
    c.has_title = False
    c.font.size = Pt(14)
    c.font.name = FONT
    c.font.color.rgb = _rgb(th["text"])
    palette = [th["accent"], th["accent2"]] + CHART_EXTRA
    plot = c.plots[0]
    if ch["type"] == "pie":
        c.has_legend = True
        c.legend.position = XL_LEGEND_POSITION.RIGHT
        c.legend.include_in_layout = False
        c.legend.font.size = Pt(14)
        c.legend.font.color.rgb = _rgb(th["text"])
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(13)
        plot.data_labels.font.bold = True
        plot.data_labels.font.color.rgb = _rgb("FFFFFF")
        plot.data_labels.number_format = "0.#"
        plot.data_labels.number_format_is_linked = False
        for i, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
    else:
        multi = len(series) > 1
        c.has_legend = multi
        if multi:
            c.legend.position = XL_LEGEND_POSITION.BOTTOM
            c.legend.include_in_layout = False
            c.legend.font.size = Pt(13)
            c.legend.font.color.rgb = _rgb(th["text"])
        for i, se in enumerate(plot.series):
            col = _rgb(palette[i % len(palette)])
            if ch["type"] == "line":
                se.format.line.color.rgb = col
                se.format.line.width = Pt(3)
                se.smooth = False
                se.marker.format.fill.solid()
                se.marker.format.fill.fore_color.rgb = col
                se.marker.format.line.color.rgb = col
            else:
                se.format.fill.solid()
                se.format.fill.fore_color.rgb = col
        if ch["type"] != "line":
            plot.gap_width = 70
        va, ca = c.value_axis, c.category_axis
        va.major_gridlines.format.line.color.rgb = _rgb(th["line"])
        va.format.line.fill.background()
        ca.format.line.color.rgb = _rgb(th["line"])
        va.tick_labels.font.size = ca.tick_labels.font.size = Pt(13)
        va.tick_labels.font.color.rgb = ca.tick_labels.font.color.rgb = _rgb(th["muted"])
    if side:
        x = MX + cw + 0.35
        w = W - MX - x
        rrect(slide, x, 2.0, w, 4.85, th["card"], line=th["line"], radius=0.05)
        rect(slide, x + 0.35, 2.35, 0.7, 0.07, th["accent"])
        _dots(slide, th, s["bullets"], x + 0.35, 2.65, w - 0.6, 3.95, 18)


# ================================================================ asosiy qurilish
def build_pptx(deck: dict, images: dict | None = None) -> bytes:
    images = images or {}
    th = THEMES.get(deck.get("theme"), THEMES[DEFAULT_THEME])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    sec_idx = 0
    dark_accents = [
        "38BDF8", "FB7185", "A78BFA", "34D399", "FBBF24", "22D3EE",
        "F472B6", "4ADE80", "C084FC", "60A5FA", "2DD4BF", "F87171",
        "A3E635", "E879F9", "818CF8", "10B981", "FDA4AF", "93C5FD",
        "D8B4FE", "5EEAD4",
    ]
    light_accents = [
        "0369A1", "BE123C", "6D28D9", "047857", "B45309", "0F766E",
        "9D174D", "166534", "7E22CE", "1D4ED8", "115E59", "B91C1C",
        "3F6212", "A21CAF", "4338CA", "065F46", "9F1239", "1E40AF",
        "6B21A8", "134E4A",
    ]
    hero_pairs = [("0B1220", "1E3A5F"), ("2A1B3D", "5B2146"),
                  ("102A43", "14532D"), ("1E1B4B", "3B1D6A")]
    for num, s in enumerate(deck.get("slides", []), 1):
        slide = prs.slides.add_slide(blank)
        lay = s.get("layout", "bullets")
        img = images.get(s.get("image_prompt", ""))
        slide_theme = th.copy()
        accent_palette = dark_accents if th["dark"] else light_accents
        accent_i = (num - 1) % len(accent_palette)
        slide_theme["accent"] = accent_palette[accent_i]
        slide_theme["accent2"] = accent_palette[(accent_i + 7) % len(accent_palette)]
        slide_theme["bg"] = th["bg"] if num % 2 else th["bg2"]
        slide_theme["bg2"] = th["bg2"] if num % 2 else th["bg"]
        slide_theme["hero1"], slide_theme["hero2"] = hero_pairs[(num - 1) % len(hero_pairs)]
        variant = (num - 1) % 3
        try:
            if lay in ("title", "closing"):
                _slide_title(slide, slide_theme, s, img, deck.get("title", "Taqdimot"), num, closing=(lay == "closing"))
            elif lay == "section":
                sec_idx += 1
                _slide_section(slide, slide_theme, s, num, sec_idx)
            else:
                _solid_bg(slide, slide_theme["bg"])
                if lay == "quote":
                    _slide_quote(slide, slide_theme, s)
                elif lay == "image_text":
                    _slide_image_text(slide, slide_theme, s, img, variant)
                elif lay == "bullets":
                    _slide_bullets(slide, slide_theme, s, variant)
                else:
                    _header(slide, slide_theme, s.get("title", "Slayd"))
                    {"bullets": _slide_bullets, "two_column": _slide_two_column, "stats": _slide_stats,
                     "steps": _slide_steps, "table": _slide_table, "chart": _slide_chart}[lay](slide, slide_theme, s)
                if not (lay == "image_text" and img):
                    _footer(slide, slide_theme, deck.get("title", "Taqdimot"), num)
        except Exception as exc:  # noqa: BLE001
            import logging
            logging.getLogger(__name__).warning("Slayd %s xatosi: %s", num, exc, exc_info=True)
        if s.get("notes"):
            try:
                slide.notes_slide.notes_text_frame.text = s["notes"]
            except Exception:  # noqa: BLE001
                pass
    prs.core_properties.title = deck.get("title", "Taqdimot")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
