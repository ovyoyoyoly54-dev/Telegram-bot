"""Regression tests for presentation request routing and PPTX delivery."""

import io
import os
import unittest
from unittest.mock import patch

os.environ.setdefault("TELEGRAM_TOKEN", "test-token")
os.environ.setdefault("GROQ_API_KEY", "test-key")

from pptx import Presentation

import bot
import config
import docs
import presenton


def _sample_pptx(slide_count=12):
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for number in range(slide_count):
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(0, 0, 100, 100)
        box.text = f"Mazmunli slayd {number + 1}"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


class PresentationRoutingTests(unittest.TestCase):
    def test_uploaded_file_summary_request_includes_document_context(self):
        message = type("Message", (), {
            "from_user": type("User", (), {"id": 42})(),
            "chat": type("Chat", (), {"id": 42})(),
        })()
        document = {"name": "report.pdf", "text": "Private report contents"}

        with (
            patch.object(docs, "get", return_value=document),
            patch.object(docs, "context_block", return_value="\nDOCUMENT CONTEXT"),
            patch.object(bot.state, "get_history", return_value=[]),
            patch.object(bot.bot, "send_chat_action"),
            patch.object(bot.llm, "chat", return_value="summary") as chat,
            patch.object(bot, "send_long"),
        ):
            bot.do_chat(
                message,
                "Men «report.pdf» faylini yubordim. Mazmunini 3 gapda tushuntir.",
            )

        self.assertIn("DOCUMENT CONTEXT", chat.call_args.args[2])

    def test_image_correction_regenerates_instead_of_answering_in_chat(self):
        correction = "Sher emas, mushuk deb qaytib yozsam"
        message = type("Message", (), {
            "from_user": type("User", (), {"id": 42})(),
            "chat": type("Chat", (), {"id": 42})(),
            "message_id": 101,
        })()

        with (
            patch.object(config, "IMAGES_ENABLED", True),
            patch.object(bot.state, "get_image", return_value=b"last image"),
            patch.object(bot.state, "get_image_prompt", return_value="Mushuk rasmini yarat"),
            patch.object(docs, "get", return_value=None),
            patch.object(bot, "_refine_image_correction", return_value="photo of a cat, not a lion"),
            patch.object(bot, "do_draw") as draw,
            patch.object(bot, "do_chat") as chat,
        ):
            bot.route_text(message, correction)

        draw.assert_called_once_with(
            message,
            "photo of a cat, not a lion",
            correction,
        )
        chat.assert_not_called()

    def test_correction_subject_after_emas_overrides_rejected_subject(self):
        captured = {}

        def refine(text):
            captured["text"] = text
            return "A cat, not a lion"

        with patch.object(bot.llm, "refine_prompt", side_effect=refine):
            bot._refine_image_correction(
                "Mushukni rasmini yarat",
                "Sher emas, mushuk deb qaytib yozsam",
            )

        self.assertIn("rejected subject is sher", captured["text"])
        self.assertIn("desired subject is mushuk", captured["text"])
        self.assertIn("do not depict sher", captured["text"])

    def test_animal_photo_request_goes_to_image_generation_not_chat(self):
        request = "Mushukni rasmini qil"
        message = type("Message", (), {
            "from_user": type("User", (), {"id": 42})(),
            "chat": type("Chat", (), {"id": 42})(),
            "message_id": 100,
        })()

        with (
            patch.object(config, "IMAGES_ENABLED", True),
            patch.object(bot.state, "get_image", return_value=None),
            patch.object(docs, "get", return_value=None),
            patch.object(bot, "_refine_image_prompt", return_value="photorealistic cat"),
            patch.object(bot, "do_draw") as draw,
            patch.object(bot, "do_chat") as chat,
            patch.object(bot.llm, "classify_vote") as classify,
        ):
            bot.route_text(message, request)

        draw.assert_called_once_with(message, "photorealistic cat", request)
        chat.assert_not_called()
        classify.assert_not_called()

    def test_uzbek_12_slide_request_generates_and_sends_real_pptx(self):
        request = "Menga sun\u2019iy intellekt haqida 12 slaydli prezentatsiya qilib ber"
        pptx_bytes = _sample_pptx(12)
        result = docs.GenerateResult(
            name="suniy_intellekt.pptx",
            data=pptx_bytes,
            text="taqdimot",
            summary="Tayyor",
        )
        message = type("Message", (), {
            "from_user": type("User", (), {"id": 42})(),
            "chat": type("Chat", (), {"id": 42})(),
            "message_id": 100,
        })()

        with (
            patch.object(config, "IMAGES_ENABLED", False),
            patch.object(config, "FILE_DAILY_LIMIT", 10),
            patch.object(bot.state, "files_left", return_value=5),
            patch.object(bot.state, "get_history", return_value=[]),
            patch.object(bot.state, "count_file"),
            patch.object(bot.state, "remember"),
            patch.object(docs, "get", return_value=None),
            patch.object(docs, "store"),
            patch.object(docs, "generate", return_value=result) as generate,
            patch.object(bot.bot, "send_chat_action"),
            patch.object(bot.bot, "send_document") as send_document,
            patch.object(bot.llm, "classify_vote") as classify,
        ):
            bot.route_text(message, request)

        generate.assert_called_once()
        self.assertEqual(generate.call_args.args[1], request)
        classify.assert_not_called()
        send_document.assert_called_once()
        sent_file = send_document.call_args.args[1]
        sent_name = sent_file.name
        sent_data = sent_file.getvalue()
        self.assertTrue(sent_name.endswith(".pptx"))
        self.assertEqual(len(Presentation(io.BytesIO(sent_data)).slides), 12)

    def test_presenton_error_falls_back_to_internal_deck(self):
        internal_result = docs.GenerateResult(
            name="fallback.pptx",
            data=_sample_pptx(12),
            text="fallback",
        )
        with (
            patch.object(config, "PRESENTON_ENABLED", True),
            patch.object(
                presenton,
                "generate_presentation",
                side_effect=presenton.PresentonError("service error"),
            ) as api,
            patch.object(docs, "_generate_deck", return_value=internal_result) as fallback,
        ):
            result = docs.generate(
                [],
                "Menga sun\u2019iy intellekt haqida 12 slaydli prezentatsiya qilib ber",
            )

        self.assertEqual(result, internal_result)
        self.assertEqual(api.call_args.kwargs["n_slides"], 12)
        fallback.assert_called_once()

    def test_presenton_uses_configured_base_url_for_generation_and_download(self):
        class Response:
            status_code = 200
            content = b"pptx-data"

            def json(self):
                return {"path": "generated/deck.pptx", "title": "Deck"}

            def raise_for_status(self):
                return None

        with (
            patch.object(config, "PRESENTON_ENABLED", True),
            patch.object(config, "PRESENTON_API_KEY", "test-presenton-key"),
            patch.object(config, "PRESENTON_BASE_URL", "https://presenton.example/api-root"),
            patch.object(presenton.requests, "post", return_value=Response()) as post,
            patch.object(presenton.requests, "get", return_value=Response()) as get,
        ):
            name, data, _ = presenton.generate_presentation("Test topic")

        self.assertEqual(post.call_args.args[0], "https://presenton.example/api-root/api/v3/presentation/generate")
        self.assertEqual(get.call_args.args[0], "https://presenton.example/api-root/generated/deck.pptx")
        self.assertEqual(get.call_args.kwargs["headers"]["Authorization"], "Bearer test-presenton-key")
        self.assertEqual(name, "Deck.pptx")
        self.assertEqual(data, b"pptx-data")

    def test_presentation_requests_are_recognized_in_three_languages(self):
        for text in (
            "Prezentatsiya qilib ber",
            "12 slaydli taqdimot yarat",
            "PPTX tayyorla",
            "\u0421\u0434\u0435\u043b\u0430\u0439 \u043f\u0440\u0435\u0437\u0435\u043d\u0442\u0430\u0446\u0438\u044e",
            "\u0421\u043e\u0437\u0434\u0430\u0439 12 \u0441\u043b\u0430\u0439\u0434\u043e\u0432",
            "Create a presentation",
            "Prepare 12 slides",
        ):
            with self.subTest(text=text):
                self.assertTrue(bot.is_presentation_request(text))

    def test_animal_photo_requests_are_recognized_in_three_languages(self):
        for text in (
            "Mushukni rasmini qil",
            "Itning real hayotdagi rasmini yubor",
            "\u041f\u0440\u0438\u0448\u043b\u0438 \u0444\u043e\u0442\u043e \u043a\u043e\u0448\u043a\u0438",
            "Send me a real photo of a cat",
        ):
            with self.subTest(text=text):
                self.assertTrue(bot.looks_like_image_request(text))

    def test_default_image_prompt_requests_a_lifelike_photo(self):
        with patch.object(bot.llm, "refine_prompt", return_value="A cat in a garden"):
            prompt = bot._refine_image_prompt("Mushukni rasmini qil")
        self.assertIn("Photorealistic real-life photograph", prompt)
        self.assertIn("No text", prompt)

    def test_explicit_art_style_is_preserved(self):
        with patch.object(bot.llm, "refine_prompt", return_value="A cartoon cat"):
            prompt = bot._refine_image_prompt("Multfilm uslubida mushuk rasmi qil")
        self.assertNotIn("Photorealistic real-life photograph", prompt)


if __name__ == "__main__":
    unittest.main()
