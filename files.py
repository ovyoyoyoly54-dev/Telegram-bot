"""Fayllarni o'qish (matn ajratish) va yaratish (docx / xlsx / pdf / pptx / oddiy matn)."""
import csv
import html
import io
import logging
import os
import re
import textwrap

import config

log = logging.getLogger("files")


class FileError(Exception):
    """Foydalanuvchiga ko'rsatish mumkin bo'lgan xato."""


BUILD_EXTS = {"docx", "xlsx", "pdf", "pptx"}
OLD_EXTS = {"doc", "xls", "ppt", "rtf", "odt", "ods", "odp"}
TEXT_EXTS = {
    "txt", "md", "markdown", "csv", "tsv", "json", "jsonl", "xml", "yaml", "yml", "toml",
    "ini", "cfg", "conf", "log", "py", "js", "ts", "jsx", "tsx", "html", "htm", "css",
    "scss", "sql", "sh", "bat", "ps1", "java", "c", "cpp", "h", "hpp", "cs", "go", "rs",
    "rb", "php", "kt", "swift", "lua", "r", "dart", "vue", "svelte", "tex", "rst", "srt",
    "ipynb", "gitignore", "dockerfile", "env.example",
}
SUPPORTED_HINT = "PDF, DOCX, XLSX, PPTX, CSV, TXT, MD, JSON, HTML va kod fayllari"


# ---------------------------------------------------------------- nomlar
def ext_of(name: str) -> str:
    return os.path.splitext(name or "")[1].lstrip(".").lower()


def safe_name(name: str, default: str = "fayl.txt") -> str:
    name = os.path.basename((name or "").replace("\\", "/")).strip()
    name = re.sub(r"[^\w.\- ]", "_", name, flags=re.U).strip(" .")
    return name[:80] or default


def clip(text: str) -> tuple[str, bool]:
    limit = config.FILE_MAX_CHARS
    if len(text) > limit:
        return text[:limit], True
    return text, False


# ================================================================ O'QISH
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise FileError("PDF parol bilan himoyalangan.")
        except FileError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise FileError("PDF parol bilan himoyalangan.") from exc
    out, total = [], 0
    for i, page in enumerate(reader.pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            out.append(f"--- {i}-sahifa ---\n{t}")
            total += len(t)
        if total > config.FILE_MAX_CHARS * 2:
            break
    return "\n\n".join(out)


def _read_docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    out, total = [], 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith("Heading"):
                m = re.search(r"\d+", style)
                lvl = min(int(m.group()) if m else 1, 6)
                out.append("#" * max(lvl, 1) + " " + t)
            elif style == "Title":
                out.append("# " + t)
            elif "List Bullet" in style:
                out.append("- " + t)
            elif "List Number" in style:
                out.append("1. " + t)
            else:
                out.append(t)
            total += len(t)
        elif tag == "tbl":
            rows = []
            for row in Table(child, doc).rows:
                cells = [c.text.strip().replace("\n", " ").replace("|", "/") for c in row.cells]
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                ncols = rows[0].count("|") - 1
                rows.insert(1, "|" + "---|" * max(ncols, 1))
                out.extend(rows)
                total += sum(len(r) for r in rows)
        if total > config.FILE_MAX_CHARS * 2:
            break
    return "\n".join(out)


def _read_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    parts = []
    try:
        for ws in wb.worksheets:
            rows = []
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= config.XLSX_MAX_ROWS:
                    break
                vals = ["" if v is None else str(v) for v in row]
                while vals and vals[-1] == "":
                    vals.pop()
                rows.append(vals)
            while rows and not rows[-1]:
                rows.pop()
            buf = io.StringIO()
            csv.writer(buf, lineterminator="\n").writerows(rows)
            parts.append(f"### {ws.title}\n{buf.getvalue()}")
    finally:
        wb.close()
    return "\n".join(parts)


def _read_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        out.append(f"# {title or f'Slayd {i}'}")
        for sh in slide.shapes:
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        out.append("  " * min(para.level, 2) + "- " + t)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    out.append("- " + " | ".join(c.text.strip() for c in row.cells))
        out.append("")
    return "\n".join(out)


def extract_text(data: bytes, name: str) -> str:
    """Fayldan matn ajratadi. Qo'llab-quvvatlanmasa yoki o'qib bo'lmasa FileError."""
    ext = ext_of(name)
    try:
        if ext == "pdf":
            text = _read_pdf(data)
        elif ext == "docx":
            text = _read_docx(data)
        elif ext in ("xlsx", "xlsm"):
            text = _read_xlsx(data)
        elif ext == "pptx":
            text = _read_pptx(data)
        elif ext in OLD_EXTS:
            raise FileError(
                f".{ext} eski format. Uni .docx / .xlsx / .pptx / .pdf ga o'tkazib qayta yuboring."
            )
        else:
            if ext not in TEXT_EXTS and b"\x00" in data[:4096]:
                raise FileError(
                    f"Bu formatni o'qiy olmayman (.{ext or '?'}). Qo'llab-quvvatlanadi: {SUPPORTED_HINT}."
                )
            text = _decode(data)
    except FileError:
        raise
    except Exception as exc:  # noqa: BLE001
        log.warning("Faylni o'qishda xato (%s): %s", name, exc)
        raise FileError("Faylni o'qib bo'lmadi: u buzilgan yoki himoyalangan bo'lishi mumkin.") from exc
    text = text.strip()
    if not text:
        hint = " Skanerlangan (rasmli) PDF bo'lishi mumkin." if ext == "pdf" else ""
        raise FileError("Faylda o'qiladigan matn topilmadi." + hint)
    return text


# ================================================================ YARATISH
_TABLE_SEP = re.compile(r"^:?-{2,}:?$")
_HEAD = re.compile(r"^(#{1,6})\s+(.*)$")
_BUL = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMB = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_RULE = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")


def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.count("|") >= 2


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _table_rows(raw_rows: list[str]) -> list[list[str]]:
    rows = [_split_row(r) for r in raw_rows]
    rows = [r for r in rows if not (any(r) and all(_TABLE_SEP.match(c) for c in r if c))]
    n = max((len(r) for r in rows), default=0)
    return [r + [""] * (n - len(r)) for r in rows]


# ---------------------------------------------------------------- docx
def _docx_runs(par, text: str) -> None:
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*(?![\s*])[^*]+?\*(?![\w*]))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            par.add_run(part[1:-1]).font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def md_to_docx(md: str) -> bytes:
    from docx import Document

    doc = Document()
    lines = md.replace("\r\n", "\n").split("\n")
    i, in_code, code = 0, False, []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                doc.add_paragraph().add_run("\n".join(code)).font.name = "Consolas"
                code, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code.append(line)
            i += 1
            continue
        if _is_table_row(line):
            raw = []
            while i < len(lines) and _is_table_row(lines[i]):
                raw.append(lines[i])
                i += 1
            rows = _table_rows(raw)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        par = table.cell(r, c).paragraphs[0]
                        _docx_runs(par, val)
                        if r == 0:
                            for run in par.runs:
                                run.bold = True
                doc.add_paragraph()
            continue
        i += 1
        if not line.strip() or _RULE.match(line):
            continue
        m = _HEAD.match(line)
        if m:
            doc.add_heading(re.sub(r"[*`]", "", m.group(2)), level=min(len(m.group(1)), 4))
            continue
        m = _BUL.match(line)
        if m:
            _docx_runs(doc.add_paragraph(style="List Bullet"), m.group(2))
            continue
        m = _NUMB.match(line)
        if m:
            _docx_runs(doc.add_paragraph(style="List Number"), m.group(1))
            continue
        _docx_runs(doc.add_paragraph(), line.strip())
    if in_code and code:
        doc.add_paragraph().add_run("\n".join(code)).font.name = "Consolas"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- xlsx
_NUM = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


def _cell_value(val: str):
    v = val.strip()
    if _NUM.match(v):
        return float(v) if "." in v else int(v)
    return val


def csv_to_xlsx(content: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    sections, name, cur = [], "Varaq1", []
    for line in content.replace("\r\n", "\n").split("\n"):
        m = re.match(r"^###\s+(.+?)\s*$", line)
        if m:
            if any(x.strip() for x in cur):
                sections.append((name, cur))
            name, cur = m.group(1), []
        else:
            cur.append(line)
    if any(x.strip() for x in cur) or not sections:
        sections.append((name, cur))

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for name, lines in sections:
        while lines and not lines[-1].strip():
            lines.pop()
        title = re.sub(r"[\[\]:*?/\\]", "", name).strip()[:31] or "Varaq"
        base, k = title, 2
        while title.lower() in used:
            suffix = f"_{k}"
            title = base[: 31 - len(suffix)] + suffix
            k += 1
        used.add(title.lower())
        ws = wb.create_sheet(title)
        rows = list(csv.reader(lines))
        widths: dict[int, int] = {}
        for r, row in enumerate(rows, 1):
            for c, val in enumerate(row, 1):
                cell = ws.cell(r, c, _cell_value(val))
                if r == 1:
                    cell.font = Font(bold=True)
                widths[c] = max(widths.get(c, 8), min(len(str(val)) + 2, 50))
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = w
        if len(rows) > 1:
            ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- pdf
_FONTS: tuple[str, str, str] | None = None


def _pdf_fonts() -> tuple[str, str, str]:
    global _FONTS
    if _FONTS:
        return _FONTS
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    here = os.path.dirname(os.path.abspath(__file__))
    for d in ("/usr/share/fonts/truetype/dejavu", os.path.join(here, "fonts")):
        reg = os.path.join(d, "DejaVuSans.ttf")
        if not os.path.exists(reg):
            continue
        bold = os.path.join(d, "DejaVuSans-Bold.ttf")
        mono = os.path.join(d, "DejaVuSansMono.ttf")
        pdfmetrics.registerFont(TTFont("DejaVu", reg))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", bold if os.path.exists(bold) else reg))
        pdfmetrics.registerFont(TTFont("DejaVu-Mono", mono if os.path.exists(mono) else reg))
        pdfmetrics.registerFontFamily(
            "DejaVu", normal="DejaVu", bold="DejaVu-Bold", italic="DejaVu", boldItalic="DejaVu-Bold"
        )
        _FONTS = ("DejaVu", "DejaVu-Bold", "DejaVu-Mono")
        return _FONTS
    log.warning("DejaVu shrifti topilmadi: PDF'da o'zbekcha/kirillcha harflar buzilishi mumkin")
    _FONTS = ("Helvetica", "Helvetica-Bold", "Courier")
    return _FONTS


def _rl_inline(text: str, mono: str) -> str:
    out = []
    for i, piece in enumerate(re.split(r"(`[^`]+`)", text)):
        if i % 2 == 1:
            out.append(f'<font name="{mono}">{html.escape(piece[1:-1], quote=False)}</font>')
        else:
            s = html.escape(piece, quote=False)
            s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
            s = re.sub(r"(?<![\w*])\*(?![\s*])(.+?)(?<![\s*])\*(?![\w*])", r"<i>\1</i>", s)
            out.append(s)
    return "".join(out)


def _wrap_code(code: str, width: int = 95) -> str:
    lines = []
    for line in code.split("\n"):
        lines.extend(textwrap.wrap(line, width, drop_whitespace=False, replace_whitespace=False) or [""])
    return "\n".join(lines)


def md_to_pdf(md: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

    reg, bold, mono = _pdf_fonts()
    body = ParagraphStyle("body", fontName=reg, fontSize=10.5, leading=15, spaceAfter=4)
    cell = ParagraphStyle("cell", parent=body, fontSize=9, leading=12, spaceAfter=0)
    heads = {
        n: ParagraphStyle(f"h{n}", parent=body, fontName=bold, fontSize=s, leading=s + 5,
                          spaceBefore=8, spaceAfter=4)
        for n, s in ((1, 20), (2, 16), (3, 13), (4, 11.5))
    }
    codest = ParagraphStyle("code", fontName=mono, fontSize=8.5, leading=11,
                            backColor=colors.whitesmoke, borderPadding=4, spaceAfter=8)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=16, bulletIndent=4)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    avail = A4[0] - 40 * mm
    story = []
    lines = md.replace("\r\n", "\n").split("\n")
    i, in_code, code = 0, False, []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                story.append(Preformatted(_wrap_code("\n".join(code)), codest))
                code, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code.append(line)
            i += 1
            continue
        if _is_table_row(line):
            raw = []
            while i < len(lines) and _is_table_row(lines[i]):
                raw.append(lines[i])
                i += 1
            rows = _table_rows(raw)
            if rows:
                n = len(rows[0])
                data = [[Paragraph(_rl_inline(c, mono), cell) for c in r] for r in rows]
                t = Table(data, colWidths=[avail / n] * n, repeatRows=1)
                t.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.extend([t, Spacer(1, 8)])
            continue
        i += 1
        if not line.strip():
            story.append(Spacer(1, 4))
            continue
        if _RULE.match(line):
            story.append(Spacer(1, 8))
            continue
        m = _HEAD.match(line)
        if m:
            story.append(Paragraph(_rl_inline(m.group(2), mono), heads[min(len(m.group(1)), 4)]))
            continue
        m = _BUL.match(line)
        if m:
            story.append(Paragraph(_rl_inline(m.group(2), mono), bullet, bulletText="•"))
            continue
        m = _NUMB.match(line)
        if m:
            num = re.match(r"^\s*(\d+)[.)]", line).group(1)
            story.append(Paragraph(_rl_inline(m.group(1), mono), bullet, bulletText=f"{num}."))
            continue
        story.append(Paragraph(_rl_inline(line.strip(), mono), body))
    if in_code and code:
        story.append(Preformatted(_wrap_code("\n".join(code)), codest))
    if not story:
        story.append(Paragraph(" ", body))
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------- pptx
def _plain(s: str) -> str:
    return re.sub(r"\*\*|__|`", "", s).strip()


def md_to_pptx(md: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt

    slides: list[tuple[str, list[tuple[int, str]]]] = []
    cur = None
    for line in md.replace("\r\n", "\n").split("\n"):
        s = line.rstrip()
        if not s.strip() or s.strip().startswith("```") or _RULE.match(s):
            continue
        m = re.match(r"^#\s+(.*)$", s)
        if m:
            cur = (_plain(m.group(1)), [])
            slides.append(cur)
            continue
        if cur is None:
            cur = ("", [])
            slides.append(cur)
        m = re.match(r"^(\s*)(?:[-*+]|\d+[.)]|#{2,6})\s+(.*)$", s)
        if m:
            cur[1].append((min(len(m.group(1)) // 2, 2), _plain(m.group(2))))
        else:
            cur[1].append((0, _plain(s)))
    if not slides:
        slides = [("Taqdimot", [])]

    prs = Presentation()
    for title, items in slides:
        chunks = [items[k:k + 7] for k in range(0, len(items), 7)] or [[]]
        for n, chunk in enumerate(chunks):
            t = (title or "Taqdimot") + (" (davomi)" if n else "")
            if chunk:
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                tf = slide.placeholders[1].text_frame
                for j, (lvl, text) in enumerate(chunk):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    p.text = text
                    p.level = lvl
                    p.font.size = Pt(22 if lvl == 0 else 18)
            else:
                slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = t
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- umumiy
def build_file(name: str, content: str) -> bytes:
    ext = ext_of(name)
    if ext == "docx":
        return md_to_docx(content)
    if ext == "xlsx":
        return csv_to_xlsx(content)
    if ext == "pdf":
        return md_to_pdf(content)
    if ext == "pptx":
        return md_to_pptx(content)
    if ext == "csv":
        return content.encode("utf-8-sig")  # Excel o'zbekcha/kirillcha harflarni to'g'ri ochsin
    return content.encode("utf-8")


def _unfence(body: str) -> str:
    m = re.match(r"^\s*```[^\n]*\n(.*)\n```\s*$", body, re.S)
    return m.group(1) if m else body


def parse_file_response(raw: str, default_name: str) -> tuple[str, str, str]:
    """Model javobi: FILENAME / SUMMARY / ===== / mazmun. (nom, mazmun, izoh) qaytaradi."""
    lines = raw.strip().split("\n")
    name = summary = ""
    idx = 0
    while idx < len(lines):
        s = lines[idx].strip()
        m = re.match(r"^(FILENAME|SUMMARY)\s*:\s*(.*)$", s, re.I)
        if m:
            if m.group(1).upper() == "FILENAME":
                name = m.group(2).strip().strip("`*\"'")
            else:
                summary = m.group(2).strip()
            idx += 1
        elif not s:
            idx += 1
        elif re.fullmatch(r"={3,}", s):
            idx += 1
            break
        else:
            break
    body = _unfence("\n".join(lines[idx:])).strip("\n")
    if not body.strip():
        raise FileError("Model fayl mazmunini qaytarmadi, qayta urinib ko'ring.")
    fallback = safe_name(default_name or "fayl.txt")
    name = safe_name(name, fallback)
    if not ext_of(name):
        name += "." + (ext_of(fallback) or "txt")
    return name, body, summary[:300]
