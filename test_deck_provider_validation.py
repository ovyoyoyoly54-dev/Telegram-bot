"""Provider independent presentation validation and failure diagnostics tests."""

import io
import json
import os
import unittest
from unittest.mock import patch

os.environ.setdefault("TELEGRAM_TOKEN", "test-token")
os.environ.setdefault("GROQ_API_KEY", "test-groq-key")

from pptx import Presentation

import config
import deck
import docs
import llm
import presenton


def _deck_payload():
    slides = [{"layout": "title", "title": "Sun'iy intellekt", "subtitle": "Mavzu va maqsad"}]
    for i in range(1, 7):
        slides.append({
            "layout": "bullets",
            "title": f"Asosiy jihat {i}",
            "bullets": [
                f"Sun'iy intellekt katta hajmdagi ma'lumotlardan andozalarni topadi va "
                f"mutaxassislarga kundalik qarorlarni izchil qabul qilishda yordam beradi, misol {i}."
            ],
        })
    slides.append({"layout": "closing", "title": "Xulosa", "subtitle": "Asosiy fikrlar va amaliy xulosa"})
    return {"title": "Sun'iy intellekt", "theme": "midnight", "slides": slides}


def _pptx(slides=8):
    presentation = Presentation()
    layout = presentation.slide_layouts[6]
    for i in range(slides):
        slide = presentation.slides.add_slide(layout)
        shape = slide.shapes.add_textbox(0, 0, 100, 100)
        shape.text = f"Slayd {i + 1} mazmuni"
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


class _Response:
    def __init__(self, status, body):
        self.status_code = status
        self.body = body
        self.text = json.dumps(body)

    def json(self):
        return self.body


class DeckValidationTests(unittest.TestCase):
    def test_json_parse_and_schema_errors_have_safe_diagnostics(self):
        with self.assertRaises(deck.DeckError):
            deck.parse_json('{"slides": [broken json]}')
        with self.assertRaisesRegex(deck.DeckError, "slides.*ro'yxati"):
            deck.normalize({"slides": "not-a-list"})
        malformed = _deck_payload()
        malformed["slides"][2]["bullets"] = {"private": "wrong type"}
        with self.assertRaisesRegex(deck.DeckError, "bullets"):
            deck.normalize(malformed)

    def test_invalid_slide_quality_is_rejected(self):
        poor = _deck_payload()
        poor["slides"][3]["bullets"] = ["Qisqa fikr."]
        normalized = deck.normalize(poor)
        self.assertTrue(docs._deck_quality_issues(normalized))

    def test_openrouter_candidate_survives_groq_and_gemini_failures(self):
        valid_json = json.dumps(_deck_payload(), ensure_ascii=False)
        groq_attempts = {"count": 0}

        def call_one(provider, *_args, **_kwargs):
            if provider == "groq":
                groq_attempts["count"] += 1
                if groq_attempts["count"] == 1:
                    return "{broken json"
                wrong_count = _deck_payload()
                wrong_count["slides"].pop()
                return json.dumps(wrong_count, ensure_ascii=False)
            if provider == "openrouter":
                return valid_json
            raise llm.LLMError("Gemini transient error", kind="server_error", status_code=503,
                               model="gemini-test", detail="provider unavailable")

        fallback = docs.GenerateResult("internal.pptx", _pptx(), "internal")
        with patch.object(config, "PRESENTATION_MAX_TOKENS", 1200), \
                patch.object(config, "IMAGES_ENABLED", False), \
                patch.object(docs, "BEST_OF_N", 2), \
                patch.object(llm, "_provider_order", return_value=["groq", "openrouter", "gemini"]), \
                patch.object(llm, "_call_one", side_effect=call_one), \
                patch.object(deck, "build_pptx", return_value=_pptx()), \
                self.assertLogs("docs", level="ERROR") as captured:
            result = docs._generate_deck([], "8 slaydli taqdimot", None, False)
        self.assertEqual(result.name, "Sun_iy intellekt.pptx")
        self.assertEqual(len(Presentation(io.BytesIO(result.data)).slides), 8)
        self.assertEqual(result.text, deck.compact(deck.normalize(_deck_payload())))
        diagnostics = "\n".join(captured.output)
        self.assertIn("error_type=DeckError", diagnostics)
        self.assertIn("error_type=ValueError", diagnostics)

    def test_all_provider_failures_keep_failure_detail(self):
        with patch.object(config, "PRESENTATION_MAX_TOKENS", 1000), \
                patch.object(llm, "_provider_order", return_value=["groq", "openrouter", "gemini"]), \
                patch.object(llm, "_call_one", side_effect=llm.LLMError("model unavailable", kind="model_not_found", status_code=404)):
            with self.assertRaisesRegex(docs.DocError, "model unavailable"):
                docs._generate_deck([], "8 slaydli taqdimot", None, False)

    def test_gemini_api_failure_keeps_http_status_model_and_safe_detail(self):
        response = _Response(500, {"error": {"code": 500, "status": "INTERNAL", "message": "backend unavailable"}})
        with patch.object(config, "GEMINI_API_KEY", "AIza-test-secret"), \
                patch.object(config, "GEMINI_MODEL", "gemini-test-model"), \
                patch.object(llm, "_post", return_value=response):
            with self.assertRaises(llm.ProviderAPIError) as caught:
                llm._chat_gemini([{"role": "user", "content": "hello"}], "system")
        self.assertEqual(caught.exception.status_code, 500)
        self.assertEqual(caught.exception.kind, "server_error")
        self.assertEqual(caught.exception.model, "gemini-test-model")
        self.assertIn("backend unavailable", caught.exception.detail)
        self.assertNotIn("AIza-test-secret", caught.exception.detail)

    def test_gemini_quota_auth_and_unknown_model_are_distinguished(self):
        quota = _Response(429, {"error": {"status": "RESOURCE_EXHAUSTED", "message": "Quota exceeded"}})
        with self.assertRaises(llm.ProviderAPIError) as caught_quota:
            llm._request_json(quota, "Gemini", model="gemini-test")
        self.assertEqual(caught_quota.exception.kind, "quota")
        self.assertEqual(caught_quota.exception.status_code, 429)
        denied = _Response(403, {"error": {"message": "API key rejected"}})
        with self.assertRaises(llm.ProviderAPIError) as caught_auth:
            llm._request_json(denied, "Gemini", model="gemini-test")
        self.assertEqual(caught_auth.exception.kind, "authentication")
        missing = _Response(404, {"error": {"message": "model not found"}})
        with self.assertRaises(llm.ProviderAPIError) as caught_model:
            llm._request_json(missing, "Gemini", model="gemini-missing")
        self.assertEqual(caught_model.exception.kind, "model_not_found")
        self.assertIn("gemini-missing", str(caught_model.exception))

    def test_presenton_500_response_detail_is_logged_and_fallback_continues(self):
        response = _Response(500, {"detail": "Presentation generation failed: worker unavailable"})
        fallback = docs.GenerateResult("internal.pptx", _pptx(), "internal")
        with patch.object(config, "PRESENTON_ENABLED", True), \
                patch.object(config, "PRESENTON_API_MODE", "cloud"), \
                patch.object(config, "PRESENTON_API_KEY", "sk-presenton-test-secret"), \
                patch.object(config, "PRESENTON_BASE_URL", "https://api.presenton.ai"), \
                patch.object(presenton.requests, "post", return_value=response), \
                patch.object(docs, "_generate_deck", return_value=fallback), \
                self.assertLogs("docs", level="WARNING") as captured:
            result = docs.generate([], "8 slaydli taqdimot")
        self.assertIs(result, fallback)
        joined = "\n".join(captured.output)
        self.assertIn("ServerError", joined)
        self.assertIn("500", joined)
        self.assertIn("worker unavailable", joined)
        self.assertNotIn("sk-presenton-test-secret", joined)


if __name__ == "__main__":
    unittest.main()
